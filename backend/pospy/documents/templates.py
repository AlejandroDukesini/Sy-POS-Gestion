"""Plantillas Excel para documentos comerciales (cotizaciones, órdenes y cualquier tipo nuevo).

Mapeo campo -> celda, sin modificar código:
  * Marcadores: la plantilla contiene {{campo}} en celdas de encabezado (también dentro de un texto:
    "Cotización N° {{numero}}") y {{linea.campo}} en la PRIMERA fila de productos. Los campos
    adicionales del tipo de documento se escriben {{extra.clave}}.
  * Manual: el administrador asigna una celda ("B5") a un campo; se guarda el valor original de la
    celda para detectar si la plantilla cambió.

La plantilla original nunca se modifica: se copia y se rellena la copia. Si la plantilla dejó de
coincidir con el mapeo se lanza TemplateMappingError con la lista de problemas, en lugar de
producir un documento incorrecto. Las fórmulas de la plantilla (p. ej. =SUMA(F11:F30)) se conservan
y se recalculan al abrir el archivo.

Limitaciones (openpyxl): no conserva imágenes, gráficos ni formas al guardar. Si la plantilla los
contiene, el relleno se hace con Microsoft Excel (COM); sin Excel instalado se informa el motivo.
"""
from __future__ import annotations

import html
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

from . import com_excel
from .excel_safety import ExcelError, ensure_workbook, sha256, unsupported_parts

HEADER_FIELDS: dict[str, str] = {
    "tipo_documento": "Tipo de documento",
    "numero": "Número",
    "fecha": "Fecha",
    "fecha_validez": "Válido hasta",
    "cliente_nombre": "Cliente",
    "cliente_identificacion": "NIT o cédula del cliente",
    "cliente_telefono": "Teléfono del cliente",
    "cliente_email": "Correo del cliente",
    "cliente_direccion": "Dirección del cliente",
    "observaciones": "Observaciones",
    "subtotal": "Subtotal (precio de lista)",
    "descuento": "Descuento",
    "base": "Base gravable",
    "iva": "IVA",
    "total": "Total",
    "vendedor": "Elaborado por",
    "empresa_nombre": "Empresa",
    "empresa_nit": "NIT de la empresa",
    "empresa_direccion": "Dirección de la empresa",
    "empresa_telefono": "Teléfono de la empresa",
    "empresa_email": "Correo de la empresa",
    "empresa_pie": "Texto al pie",
}
LINE_FIELDS: dict[str, str] = {
    "n": "N°", "codigo": "Código", "descripcion": "Descripción", "unidad": "Unidad", "cantidad": "Cantidad",
    "precio_lista": "Precio sin descuento", "precio": "Precio unitario", "descuento": "Descuento",
    "iva_pct": "IVA %", "total": "Total línea",
}
REQUIRED_HEADER = ("numero", "cliente_nombre", "total")
REQUIRED_LINE = ("descripcion", "cantidad")
NUMERIC = {"subtotal", "descuento", "base", "iva", "total", "linea.cantidad", "linea.precio", "linea.precio_lista",
           "linea.descuento", "linea.total", "linea.n", "linea.iva_pct"}
PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_\.]*)\s*\}\}")
PAPER = {"letter": 1, "legal": 5, "a4": 9, "a5": 11}


class TemplateMappingError(ExcelError):
    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("La plantilla no coincide con el mapeo: " + " ".join(problems), {"template": "; ".join(problems)})


@dataclass
class Mapping:
    sheet: str = ""
    fields: dict[str, str] = field(default_factory=dict)        # campo -> "B5"
    line_row: int = 0
    line_columns: dict[str, str] = field(default_factory=dict)  # campo -> "C"
    capacity: int = 0
    originals: dict[str, Any] = field(default_factory=dict)
    manual: list[str] = field(default_factory=list)
    fingerprint: str = ""

    def to_dict(self) -> dict:
        return dict(self.__dict__)

    @classmethod
    def from_dict(cls, data: dict | None) -> Mapping:
        known = {k: v for k, v in (data or {}).items() if k in cls.__dataclass_fields__}
        return cls(**known)


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _known(name: str, extra_keys: set[str]) -> bool:
    return name in HEADER_FIELDS or (name.startswith("extra.") and name[6:] in extra_keys)


def scan(path: Path, extra_keys: set[str], previous: Mapping | None = None) -> Mapping:
    """Detecta marcadores y conserva las asignaciones manuales previas que sigan siendo válidas."""
    ensure_workbook(path)
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet = previous.sheet if previous and previous.sheet in wb.sheetnames else wb.worksheets[0].title
        ws = wb[sheet]
        mapping = Mapping(sheet=sheet, fingerprint=sha256(path))
        for row in ws.iter_rows():
            for cell in row:
                for name in PLACEHOLDER.findall(_text(cell.value)):
                    if name.startswith("linea."):
                        fld = name.split(".", 1)[1]
                        if fld in LINE_FIELDS and (not mapping.line_row or cell.row == mapping.line_row):
                            mapping.line_row = cell.row
                            mapping.line_columns[fld] = cell.column_letter
                            mapping.originals[cell.coordinate] = cell.value
                    elif _known(name, extra_keys) and name not in mapping.fields:
                        mapping.fields[name] = cell.coordinate
                        mapping.originals[cell.coordinate] = cell.value
        if previous:
            for fld in previous.manual:
                coord = previous.fields.get(fld)
                if coord and fld not in mapping.fields and _known(fld, extra_keys):
                    mapping.fields[fld] = coord
                    mapping.manual.append(fld)
                    mapping.originals[coord] = ws[coord].value
            if previous.capacity and previous.line_row == mapping.line_row:
                mapping.capacity = previous.capacity
        if mapping.line_row and not mapping.capacity:
            mapping.capacity = _detect_capacity(ws, mapping)
        return mapping
    finally:
        wb.close()


def _detect_capacity(ws, mapping: Mapping) -> int:
    """Filas libres bajo la primera fila de productos, hasta encontrar contenido o un cambio de bordes."""
    cols = [column_index_from_string(c) for c in mapping.line_columns.values()]
    lo, hi = min(cols), max(cols)
    ref_col = column_index_from_string(mapping.line_columns.get("descripcion") or next(iter(mapping.line_columns.values())))

    def border_key(cell) -> tuple:
        return tuple(getattr(cell.border, side).style for side in ("left", "right", "top", "bottom"))

    ref_border = border_key(ws.cell(row=mapping.line_row, column=ref_col))
    styled = any(ref_border)
    capacity = 1
    for r in range(mapping.line_row + 1, mapping.line_row + 500):
        if r > ws.max_row:
            break
        values = [ws.cell(row=r, column=c).value for c in range(lo, hi + 1)]
        if any(v not in (None, "") and not (isinstance(v, str) and v.startswith("=")) for v in values):
            break
        if styled and border_key(ws.cell(row=r, column=ref_col)) != ref_border:
            break
        capacity += 1
    return capacity


def assign(mapping: Mapping, path: Path, field_name: str, coordinate: str, extra_keys: set[str]) -> Mapping:
    coordinate = coordinate.strip().upper().split("!")[-1]
    try:
        coordinate_to_tuple(coordinate)
    except Exception:  # noqa: BLE001
        raise ExcelError(f"«{coordinate}» no es una celda válida (ej.: B5).", {"cell": "Celda no válida."}) from None
    if not _known(field_name, extra_keys):
        raise ExcelError(f"Campo desconocido: {field_name}")
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        mapping.fields[field_name] = coordinate
        mapping.originals[coordinate] = wb[mapping.sheet][coordinate].value
        if field_name not in mapping.manual:
            mapping.manual.append(field_name)
    finally:
        wb.close()
    return mapping


def unassign(mapping: Mapping, field_name: str) -> Mapping:
    if field_name in mapping.manual:
        mapping.manual.remove(field_name)
        coord = mapping.fields.pop(field_name, None)
        if coord and coord not in mapping.fields.values():
            mapping.originals.pop(coord, None)
    return mapping


def validate(path: Path, mapping: Mapping) -> list[str]:
    problems: list[str] = []
    if not mapping.sheet:
        return ["La plantilla no ha sido analizada."]
    for fld in REQUIRED_HEADER:
        if fld not in mapping.fields:
            problems.append(f"Falta el campo obligatorio «{HEADER_FIELDS[fld]}» ({{{{{fld}}}}}).")
    if not mapping.line_row:
        problems.append("No se encontró la fila de productos ({{linea.descripcion}}, {{linea.cantidad}}…).")
    for fld in REQUIRED_LINE:
        if mapping.line_row and fld not in mapping.line_columns:
            problems.append(f"Falta la columna de productos «{LINE_FIELDS[fld]}» ({{{{linea.{fld}}}}}).")
    ensure_workbook(path)
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        if mapping.sheet not in wb.sheetnames:
            return problems + [f"La hoja «{mapping.sheet}» ya no existe en la plantilla."]
        ws = wb[mapping.sheet]
        for fld, coord in mapping.fields.items():
            current = _text(ws[coord].value)
            if fld in mapping.manual:
                if coord in mapping.originals and current != _text(mapping.originals[coord]):
                    problems.append(f"La celda {coord} asignada a «{HEADER_FIELDS.get(fld, fld)}» cambió en la plantilla.")
            elif fld not in [n for n in PLACEHOLDER.findall(current)]:
                problems.append(f"No se encuentra {{{{{fld}}}}} en {coord}.")
        for fld, col in mapping.line_columns.items():
            coord = f"{col}{mapping.line_row}"
            if f"linea.{fld}" not in PLACEHOLDER.findall(_text(ws[coord].value)):
                problems.append(f"No se encuentra {{{{linea.{fld}}}}} en {coord}.")
    finally:
        wb.close()
    return problems


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        if value.is_integer():
            return f"{value:,.0f}".replace(",", ".")
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return _text(value)


def _fill(template_value: Any, values: dict[str, Any], prefix: str = "") -> Any:
    """Sustituye marcadores. Una celda que es solo un marcador numérico recibe el número (conserva su formato)."""
    text = _text(template_value)
    only = PLACEHOLDER.fullmatch(text.strip())
    if only:
        key = only.group(1)
        value = values.get(key[len(prefix):] if prefix and key.startswith(prefix) else key, "")
        return value if key in NUMERIC else ("" if value is None else value)

    def repl(m: re.Match) -> str:
        key = m.group(1)
        return _fmt(values.get(key[len(prefix):] if prefix and key.startswith(prefix) else key, ""))
    return PLACEHOLDER.sub(repl, text)


def build_cells(mapping: Mapping, header: dict[str, Any], lines: list[dict[str, Any]], template_values: dict[str, Any]) -> dict:
    if len(lines) > mapping.capacity:
        raise TemplateMappingError([f"El documento tiene {len(lines)} líneas y la plantilla admite {mapping.capacity}. "
                                    "Amplíe la zona de productos de la plantilla o divida el documento."])
    cells: dict[str, Any] = {}
    for fld, coord in mapping.fields.items():
        cells[coord] = header.get(fld, "") if fld in mapping.manual else _fill(template_values.get(coord), header)
    for i in range(mapping.capacity):
        row = mapping.line_row + i
        line = lines[i] if i < len(lines) else None
        for col in mapping.line_columns.values():
            source = template_values.get(f"{col}{mapping.line_row}")
            cells[f"{col}{row}"] = None if line is None else _fill(source, line, prefix="linea.")
    return cells


def render(template: Path, mapping: Mapping, out: Path, header: dict, lines: list[dict], page_setup: dict) -> Path:
    template, out = Path(template), Path(out)
    if template.resolve() == out.resolve():
        raise ExcelError("La salida no puede ser la plantilla original.")
    problems = validate(template, mapping)
    if problems:
        raise TemplateMappingError(problems)
    wb = openpyxl.load_workbook(template, data_only=False)
    template_values = {c.coordinate: c.value for row in wb[mapping.sheet].iter_rows() for c in row if c.value is not None}
    wb.close()
    cells = build_cells(mapping, header, lines, template_values)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f".~tmp_{out.name}")
    shutil.copy2(template, tmp)
    try:
        if unsupported_parts(template):
            if not com_excel.com_available():
                raise ExcelError("La plantilla contiene " + ", ".join(unsupported_parts(template)) +
                                 ". Para conservarlos se necesita Microsoft Excel instalado en el equipo del servidor.")
            com_excel.write_cells(tmp, mapping.sheet, cells, page_setup)
        else:
            wb = openpyxl.load_workbook(tmp, data_only=False, keep_vba=tmp.suffix.lower() == ".xlsm")
            ws = wb[mapping.sheet]
            for coord, value in cells.items():
                ws[coord].value = value
            apply_page_setup(ws, page_setup)
            wb.calculation.fullCalcOnLoad = True
            wb.save(tmp)
            wb.close()
        tmp.replace(out)
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except PermissionError:   # Excel aún libera el archivo; se reemplaza en la próxima generación
            pass
    return out


def apply_page_setup(ws, setup: dict) -> None:
    ws.page_setup.orientation = "landscape" if setup.get("orientation") == "landscape" else "portrait"
    ws.page_setup.paperSize = PAPER.get(str(setup.get("paper", "letter")).lower(), 1)
    margin_in = float(setup.get("margins_cm", 1.5)) / 2.54
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = margin_in
    if setup.get("fit_width", True):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = max(10, min(400, int(setup.get("scale", 100))))
    if not ws.print_area:
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


# --------------------------------------------------------------------------- vista previa HTML

def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb != "00000000":
            return "#" + c.rgb[2:]
    except AttributeError:
        pass
    return None


def _display(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        return "(se calcula al abrir)"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = cell.number_format or ""
        if "$" in fmt:
            return "$ " + _fmt(float(value))
        return _fmt(float(value))
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    return str(value)


def preview_html(path: Path, sheet: str | None = None, max_rows: int = 200, max_cols: int = 30) -> str:
    """Tabla HTML aproximada de la hoja (valores, combinaciones, negrita, rellenos, bordes y anchos).

    Es una vista previa de contenido, no una reproducción exacta de la impresión de Excel.
    """
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        last_row, last_col = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
        if ws.print_area:
            try:
                area = str(ws.print_area).split("!")[-1].replace("$", "").split(":")
                last_col = min(column_index_from_string(re.sub(r"\d", "", area[1])), max_cols)
                last_row = min(int(re.sub(r"\D", "", area[1])), max_rows)
            except (ValueError, IndexError):
                pass
        spans: dict[tuple[int, int], tuple[int, int]] = {}
        hidden: set[tuple[int, int]] = set()
        for rng in ws.merged_cells.ranges:
            spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        hidden.add((r, c))
        cols = []
        for c in range(1, last_col + 1):
            width = ws.column_dimensions[get_column_letter(c)].width or 9
            cols.append(f'<col style="width:{round(width * 7.2)}px">')
        out = ['<table class="xl-preview">', "<colgroup>", *cols, "</colgroup>"]
        for r in range(1, last_row + 1):
            out.append("<tr>")
            for c in range(1, last_col + 1):
                if (r, c) in hidden:
                    continue
                cell = ws.cell(row=r, column=c)
                style = []
                if cell.font is not None:
                    if cell.font.b:
                        style.append("font-weight:700")
                    if cell.font.sz and cell.font.sz >= 13:
                        style.append(f"font-size:{round(cell.font.sz / 11, 2)}em")
                    if (fc := _color(cell.font.color)):
                        style.append(f"color:{fc}")
                if cell.fill is not None and cell.fill.fill_type == "solid" and (bg := _color(cell.fill.fgColor)):
                    style.append(f"background:{bg}")
                if cell.alignment is not None and cell.alignment.horizontal in ("center", "right", "left"):
                    style.append(f"text-align:{cell.alignment.horizontal}")
                elif isinstance(cell.value, (int, float)):
                    style.append("text-align:right")
                for side in ("left", "right", "top", "bottom"):
                    if getattr(cell.border, side).style:
                        style.append(f"border-{side}:1px solid #6b7280")
                rowspan, colspan = spans.get((r, c), (1, 1))
                attrs = (f' rowspan="{rowspan}"' if rowspan > 1 else "") + (f' colspan="{colspan}"' if colspan > 1 else "")
                out.append(f'<td{attrs} style="{";".join(style)}">{html.escape(_display(cell))}</td>')
            out.append("</tr>")
        out.append("</table>")
        return "".join(out)
    finally:
        wb.close()


# --------------------------------------------------------------------------- plantilla base

def create_sample(title: str, path: Path, extra_fields: list[dict] | None = None, lines: int = 20) -> Path:
    """Plantilla base editable en Excel, con todos los marcadores, como punto de partida."""
    path = Path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Documento"
    dark = PatternFill("solid", fgColor="1F2937")
    soft = PatternFill("solid", fgColor="EEF2F7")
    thin = Side(style="thin", color="9CA3AF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, w in {"A": 6, "B": 14, "C": 44, "D": 10, "E": 16, "F": 18}.items():
        ws.column_dimensions[col].width = w
    ws["A1"] = "{{empresa_nombre}}"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = "NIT {{empresa_nit}} · {{empresa_direccion}}"
    ws["A3"] = "Tel. {{empresa_telefono}} · {{empresa_email}}"
    ws["E1"] = title.upper()
    ws["E1"].font = Font(size=14, bold=True, color="1D4ED8")
    ws["E2"] = "N° {{numero}}"
    ws["E2"].font = Font(bold=True)
    ws["E3"] = "Fecha: {{fecha}}"
    ws["E4"] = "Válido hasta: {{fecha_validez}}"
    labels = [("A6", "Cliente:", "B6", "{{cliente_nombre}}"), ("A7", "NIT/CC:", "B7", "{{cliente_identificacion}}"),
              ("A8", "Teléfono:", "B8", "{{cliente_telefono}}"), ("D6", "Correo:", "E6", "{{cliente_email}}"),
              ("D7", "Dirección:", "E7", "{{cliente_direccion}}"), ("D8", "Elaboró:", "E8", "{{vendedor}}")]
    for lc, lt, vc, vt in labels:
        ws[lc] = lt
        ws[lc].font = Font(bold=True)
        ws[vc] = vt
    for r in (6, 7, 8):
        ws.merge_cells(f"B{r}:C{r}")
        ws.merge_cells(f"E{r}:F{r}")
    row = 9
    for extra in extra_fields or []:
        ws[f"A{row}"] = f"{extra['label']}:"
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = "{{extra.%s}}" % extra["key"]
        ws.merge_cells(f"B{row}:F{row}")
        row += 1
    head_row = row + 1
    for col, text in zip("ABCDEF", ["N°", "Código", "Descripción", "Cant.", "Precio unit.", "Total"], strict=True):
        c = ws[f"{col}{head_row}"]
        c.value = text
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = dark
        c.alignment = Alignment(horizontal="center")
        c.border = box
    first = head_row + 1
    for i in range(lines):
        r = first + i
        for col in "ABCDEF":
            ws[f"{col}{r}"].border = box
            if i % 2:
                ws[f"{col}{r}"].fill = soft
        ws[f"D{r}"].number_format = "#,##0.##"
        ws[f"E{r}"].number_format = '"$"#,##0'
        ws[f"F{r}"].number_format = '"$"#,##0'
    for col, fld in zip("ABCDEF", ["n", "codigo", "descripcion", "cantidad", "precio", "total"], strict=True):
        ws[f"{col}{first}"] = "{{linea.%s}}" % fld
    tot = first + lines + 1
    for offset, (label, fld) in enumerate([("Subtotal", "subtotal"), ("Descuento", "descuento"), ("IVA incluido", "iva"),
                                           ("TOTAL", "total")]):
        r = tot + offset
        ws[f"E{r}"] = label
        ws[f"E{r}"].font = Font(bold=True)
        ws[f"F{r}"] = "{{%s}}" % fld
        ws[f"F{r}"].number_format = '"$"#,##0'
        ws[f"F{r}"].border = box
    ws[f"A{tot}"] = "Observaciones:"
    ws[f"A{tot}"].font = Font(bold=True)
    ws[f"A{tot + 1}"] = "{{observaciones}}"
    ws.merge_cells(f"A{tot + 1}:D{tot + 3}")
    ws[f"A{tot + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws[f"A{tot + 5}"] = "{{empresa_pie}}"
    ws.merge_cells(f"A{tot + 5}:F{tot + 5}")
    ws.print_area = f"A1:F{tot + 5}"
    ws.page_setup.paperSize = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
