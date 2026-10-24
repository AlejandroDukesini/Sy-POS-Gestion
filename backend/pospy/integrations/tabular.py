"""Lectura de inventario desde Excel (.xlsx/.xlsm) o CSV, con mapeo de columnas y validación por fila."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

from ..core.errors import ValidationError
from ..core.money import parse_decimal

FIELDS = {
    "code": "Código", "description": "Descripción", "price": "Precio sin descuento",
    "discount_price": "Precio con descuento", "cost": "Costo", "stock": "Cantidad disponible",
    "min_stock": "Existencia mínima", "category": "Categoría", "tax_rate": "IVA %", "unit": "Unidad",
}
REQUIRED = ("code", "description", "price")
SYNONYMS = {
    "code": ("codigo", "cod", "referencia", "ref", "sku", "codigo producto", "item"),
    "description": ("descripcion", "nombre", "producto", "detalle", "articulo"),
    "price": ("precio", "precio venta", "precio sin descuento", "precio lista", "pvp", "valor", "precio unitario"),
    "discount_price": ("precio con descuento", "precio descuento", "precio oferta", "oferta", "precio promocion"),
    "cost": ("costo", "costo unitario", "precio compra", "costo promedio"),
    "stock": ("cantidad", "existencia", "existencias", "stock", "cantidad disponible", "saldo", "inventario"),
    "min_stock": ("minimo", "existencia minima", "stock minimo", "cantidad minima"),
    "category": ("categoria", "linea", "grupo", "familia"),
    "tax_rate": ("iva", "iva %", "impuesto", "tarifa iva"),
    "unit": ("unidad", "um", "unidad de medida"),
}
MAX_ROWS = 50_000


def _norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9% ]", " ", text)).strip()


@dataclass
class Table:
    sheets: list[str]
    sheet: str
    header_row: int
    headers: list[str]
    rows: list[list]          # filas de datos (sin encabezado)


@dataclass
class Record:
    row: int
    code: str
    values: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)


def read_table(content: bytes, filename: str, sheet: str = "", header_row: int = 0) -> Table:
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            class dialect(csv.excel):  # noqa: N801
                delimiter = ";" if text.count(";") > text.count(",") else ","
        grid = [row for row in csv.reader(io.StringIO(text), dialect)]
        sheets, sheet = ["CSV"], "CSV"
    elif name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl lanza tipos variados ante archivos dañados
            raise ValidationError(f"No se pudo leer el libro de Excel: {exc}") from None
        try:
            sheets = wb.sheetnames
            sheet = sheet if sheet in sheets else sheets[0]
            grid = []
            for i, row in enumerate(wb[sheet].iter_rows(values_only=True)):
                if i > MAX_ROWS + 50:
                    break
                grid.append(list(row))
        finally:
            wb.close()
    else:
        raise ValidationError("Formato no soportado. Use .xlsx, .xlsm o .csv.")
    if header_row <= 0:
        header_row = _detect_header(grid)
    if header_row > len(grid):
        raise ValidationError("La fila de encabezados está fuera del archivo.")
    headers = [str(h).strip() if h is not None else "" for h in grid[header_row - 1]] if grid else []
    rows = grid[header_row:]
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"El archivo tiene más de {MAX_ROWS} filas.")
    return Table(sheets, sheet, header_row, headers, rows)


def _detect_header(grid: list[list]) -> int:
    """La fila (de las primeras 20) que más coincide con nombres de columna conocidos."""
    known = {s for syn in SYNONYMS.values() for s in syn}
    best, best_row = 0, 1
    for i, row in enumerate(grid[:20], start=1):
        score = sum(1 for cell in row if _norm(cell) in known)
        if score > best:
            best, best_row = score, i
    return best_row


def suggest_columns(headers: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    normalized = {_norm(h): h for h in headers if h}
    for fld, synonyms in SYNONYMS.items():
        for syn in synonyms:
            if syn in normalized and normalized[syn] not in out.values():
                out[fld] = normalized[syn]
                break
    return out


def to_records(table: Table, columns: dict[str, str]) -> list[Record]:
    missing = [FIELDS[f] for f in REQUIRED if not columns.get(f)]
    if missing:
        raise ValidationError("Asigne las columnas obligatorias: " + ", ".join(missing), {"columns": "Faltan columnas."})
    index = {}
    for fld, header in columns.items():
        if not header:
            continue
        if fld not in FIELDS:
            raise ValidationError(f"Campo desconocido en el mapeo: {fld}")
        if header not in table.headers:
            raise ValidationError(f"La columna «{header}» no existe en el archivo.", {"columns": f"«{header}» no encontrada."})
        index[fld] = table.headers.index(header)
    records = []
    for offset, row in enumerate(table.rows):
        row_no = table.header_row + offset + 1
        cells = {fld: (row[i] if i < len(row) else None) for fld, i in index.items()}
        if all(v in (None, "") for v in cells.values()):
            continue
        rec = Record(row=row_no, code=str(cells.get("code") or "").strip())
        if isinstance(cells.get("code"), float) and cells["code"].is_integer():
            rec.code = str(int(cells["code"]))
        if not rec.code:
            rec.errors.append("Sin código.")
        desc = str(cells.get("description") or "").strip()
        if not desc:
            rec.errors.append("Sin descripción.")
        rec.values["description"] = desc[:200]
        for fld in ("price", "discount_price", "cost", "stock", "min_stock", "tax_rate"):
            if fld not in cells or cells[fld] in (None, ""):
                continue
            number = parse_decimal(cells[fld])
            if number is None:
                rec.errors.append(f"{FIELDS[fld]}: «{cells[fld]}» no es un número.")
            elif number < 0 and fld != "stock":
                rec.errors.append(f"{FIELDS[fld]} negativo.")
            else:
                if fld == "tax_rate" and Decimal(0) < number < 1:
                    number = number * 100     # 0,19 en Excel con formato de porcentaje
                rec.values[fld] = number
        if "price" not in rec.values and not any("Precio sin descuento" in e for e in rec.errors):
            rec.errors.append("Sin precio.")
        if "discount_price" in rec.values and "price" in rec.values and rec.values["discount_price"] > rec.values["price"]:
            rec.errors.append("El precio con descuento es mayor que el precio sin descuento.")
        for fld in ("category", "unit"):
            if cells.get(fld) not in (None, ""):
                rec.values[fld] = str(cells[fld]).strip()[:80]
        records.append(rec)
    return records
