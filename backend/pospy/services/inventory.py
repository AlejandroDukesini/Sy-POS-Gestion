"""Productos y movimientos de inventario.

La existencia de un producto es la suma de sus movimientos (la base de datos lo garantiza con
triggers). Editar un producto nunca cambia su existencia ni borra historial: para eso están las
entradas, salidas y ajustes, que quedan con fecha, usuario, motivo, cantidades y asiento contable.
"""
from __future__ import annotations

import csv
import io
import sqlite3
from datetime import timedelta
from decimal import Decimal

from ..core.actor import Actor
from ..core.audit import record
from ..core.clock import iso, local_today, parse_date
from ..core.errors import ConflictError, NotFoundError, ValidationError
from ..core.money import from_cents, from_milli, parse_decimal, round_cents, to_bp, to_cents, to_milli
from ..core.query import like, paginate
from . import ledger
from .ledger import Line
from .open_items import AP, OpenItems

UNITS = ("und", "kg", "g", "l", "ml", "m", "cm", "caja", "paquete", "servicio")
ENTRY_COUNTERPARTS = {
    "compra_contado": "Compra de contado",
    "compra_credito": "Compra a crédito (crea cuenta por pagar)",
    "inventario_inicial": "Inventario inicial (aporte de capital)",
    "sobrante": "Sobrante encontrado",
}
EXIT_COUNTERPARTS = {
    "merma": "Merma, daño o pérdida",
    "consumo_interno": "Consumo interno",
}
PRODUCT_FIELDS = ("code", "description", "category_id", "unit", "price", "discount_price", "cost", "tax_rate",
                  "min_stock", "active", "external_ref")


class InventoryService:
    def __init__(self, db, settings):
        self.db = db
        self.settings = settings
        self.payables = OpenItems(db, settings, AP)

    def _tz(self) -> str:
        return self.settings.get("empresa")["timezone"]

    # ------------------------------------------------------------------ consulta
    def list(self, actor: Actor, q: str = "", category_id: int | None = None, status: str = "activos",
             low_stock: bool = False, sort: str | None = None, direction: str | None = None, page: int = 1,
             page_size: int = 50) -> dict:
        actor.require("inventario.ver")
        where, params = [], []
        if status == "activos":
            where.append("p.active = 1")
        elif status == "inactivos":
            where.append("p.active = 0")
        if q:
            where.append("(p.code LIKE ? ESCAPE '\\' OR p.description LIKE ? ESCAPE '\\')")
            params += [like(q.strip())] * 2
        if category_id:
            where.append("p.category_id = ?")
            params.append(category_id)
        if low_stock:
            where.append("p.stock_milli <= p.min_stock_milli AND p.active = 1")
        sql = ("SELECT p.*, c.name AS category_name, (p.stock_milli <= p.min_stock_milli) AS low_stock"
               " FROM products p LEFT JOIN categories c ON c.id = p.category_id"
               + (" WHERE " + " AND ".join(where) if where else ""))
        with self.db.read() as conn:
            result = paginate(conn, sql, params, sort=sort, direction=direction, default_sort="code", page=page,
                              page_size=page_size, tiebreak="p.id", sortable={
                                  "code": "p.code COLLATE NOCASE", "description": "p.description COLLATE NOCASE",
                                  "price": "p.price_cents", "discount_price": "COALESCE(p.discount_price_cents, p.price_cents)",
                                  "cost": "p.cost_cents", "stock": "p.stock_milli", "updated_at": "p.updated_at"})
        result["items"] = [self._public(actor, p) for p in result["items"]]
        return result

    @staticmethod
    def _public(actor: Actor, product: dict) -> dict:
        product["low_stock"] = bool(product.get("low_stock")) and bool(product["active"])
        if not actor.can("inventario.costos"):
            product.pop("cost_cents", None)
        return product

    def get(self, actor: Actor, product_id: int) -> dict:
        actor.require("inventario.ver")
        with self.db.read() as conn:
            row = conn.execute("SELECT p.*, c.name AS category_name, (p.stock_milli <= p.min_stock_milli) AS low_stock,"
                               " cu.username AS created_by_name, uu.username AS updated_by_name FROM products p"
                               " LEFT JOIN categories c ON c.id = p.category_id LEFT JOIN users cu ON cu.id = p.created_by"
                               " LEFT JOIN users uu ON uu.id = p.updated_by WHERE p.id = ?", (product_id,)).fetchone()
        if not row:
            raise NotFoundError("Producto no encontrado.")
        return self._public(actor, dict(row))

    def lookup(self, actor: Actor, q: str, limit: int = 20) -> list[dict]:
        """Búsqueda rápida para vender o cotizar: coincidencia exacta de código primero."""
        if not any(actor.can(p) for p in ("ventas.vender", "documentos.crear", "inventario.ver")):
            actor.require("inventario.ver")
        q = (q or "").strip()
        if not q:
            return []
        with self.db.read() as conn:
            rows = conn.execute(
                "SELECT p.id, p.code, p.description, p.unit, p.price_cents, p.discount_price_cents, p.tax_rate_bp,"
                " p.stock_milli, p.min_stock_milli FROM products p WHERE p.active = 1 AND (p.code = ? COLLATE NOCASE"
                " OR p.code LIKE ? ESCAPE '\\' OR p.description LIKE ? ESCAPE '\\')"
                " ORDER BY (p.code = ? COLLATE NOCASE) DESC, p.description COLLATE NOCASE LIMIT ?",
                (q, like(q), like(q), q, min(50, max(1, limit)))).fetchall()
        return [dict(r) for r in rows]

    def history(self, actor: Actor, product_id: int | None = None, kind: str = "", date_from: str = "",
                date_to: str = "", page: int = 1, page_size: int = 50) -> dict:
        actor.require("inventario.ver")
        where, params = [], []
        if product_id:
            where.append("m.product_id = ?")
            params.append(product_id)
        if kind:
            where.append("m.kind = ?")
            params.append(kind)
        if date_from:
            where.append("m.created_at >= ?")
            params.append(date_from[:10])
        if date_to:
            where.append("m.created_at < date(?, '+1 day')")
            params.append(date_to[:10])
        sql = ("SELECT m.*, p.code, p.description, u.username, u.full_name FROM inventory_movements m"
               " JOIN products p ON p.id = m.product_id LEFT JOIN users u ON u.id = m.user_id"
               + (" WHERE " + " AND ".join(where) if where else ""))
        with self.db.read() as conn:
            result = paginate(conn, sql, params, sort="id", direction="desc", default_sort="id", tiebreak="m.id", sortable={"id": "m.id"},
                              page=page, page_size=page_size)
        if not actor.can("inventario.costos"):
            for m in result["items"]:
                m.pop("unit_cost_cents", None)
        return result

    def summary(self, actor: Actor) -> dict:
        actor.require("inventario.ver")
        with self.db.read() as conn:
            row = conn.execute("SELECT COUNT(*) AS products, SUM(stock_milli <= min_stock_milli) AS low_stock,"
                               " SUM(stock_milli <= 0) AS out_of_stock,"
                               " SUM(CASE WHEN stock_milli > 0 THEN stock_milli * cost_cents / 1000 ELSE 0 END) AS value_cents"
                               " FROM products WHERE active = 1").fetchone()
        result = {"products": row["products"], "low_stock": row["low_stock"] or 0, "out_of_stock": row["out_of_stock"] or 0}
        if actor.can("inventario.costos"):
            result["value_cents"] = int(row["value_cents"] or 0)
        return result

    # ------------------------------------------------------------------ altas y ediciones
    def _validated(self, conn: sqlite3.Connection, data: dict, current: dict | None) -> dict:
        errors: dict[str, str] = {}
        merged = {**(current or {}), **{k: v for k, v in data.items() if k in PRODUCT_FIELDS}}

        def money(key: str, label: str, required: bool) -> int | None:
            raw = merged.get(key) if key in data or current is None else None
            if key not in data and current is not None:
                return current[f"{key}_cents"]
            if raw in (None, ""):
                if required:
                    errors[key] = f"{label} es obligatorio."
                return None
            value = parse_decimal(raw)
            if value is None or value < 0:
                errors[key] = f"{label}: escriba un valor numérico mayor o igual a cero."
                return None
            if value > Decimal("100000000000"):
                errors[key] = f"{label}: valor demasiado alto."
                return None
            return to_cents(value)

        code = str(merged.get("code") or "").strip()
        description = str(merged.get("description") or "").strip()
        if not code:
            errors["code"] = "El código es obligatorio."
        elif len(code) > 40:
            errors["code"] = "El código admite hasta 40 caracteres."
        if not description:
            errors["description"] = "La descripción es obligatoria."
        elif len(description) > 200:
            errors["description"] = "La descripción admite hasta 200 caracteres."
        price = money("price", "El precio sin descuento", True)
        discount_price = money("discount_price", "El precio con descuento", False)
        cost = money("cost", "El costo", False)
        if price is not None and discount_price is not None and discount_price > price:
            errors["discount_price"] = "El precio con descuento no puede ser mayor que el precio sin descuento."
        unit = str(merged.get("unit") or "und").strip()
        if unit not in UNITS:
            errors["unit"] = "Unidad no válida."
        tax = merged.get("tax_rate") if "tax_rate" in data or current is None else None
        tax_bp = current["tax_rate_bp"] if tax is None and current is not None else None
        if tax_bp is None:
            parsed = parse_decimal(tax if tax not in (None, "") else 0)
            if parsed is None or parsed < 0 or parsed > 100:
                errors["tax_rate"] = "La tarifa de IVA debe estar entre 0 y 100 %."
            else:
                tax_bp = to_bp(parsed)
        min_stock = merged.get("min_stock") if "min_stock" in data or current is None else None
        if min_stock is None and current is not None:
            min_milli = current["min_stock_milli"]
        else:
            parsed = parse_decimal(min_stock if min_stock not in (None, "") else self.settings.get("inventario")["default_min_stock"])
            if parsed is None or parsed < 0:
                errors["min_stock"] = "La existencia mínima debe ser un número mayor o igual a cero."
                min_milli = 0
            else:
                min_milli = to_milli(parsed)
        category_id = merged.get("category_id") or None
        if category_id is not None:
            try:
                category_id = int(category_id)
            except (TypeError, ValueError):
                category_id = -1
            if not conn.execute("SELECT 1 FROM categories WHERE id = ?", (category_id,)).fetchone():
                errors["category_id"] = "Seleccione una categoría válida."
        if code and not errors.get("code"):
            dup = conn.execute("SELECT id FROM products WHERE code = ? AND id IS NOT ?",
                               (code, current["id"] if current else None)).fetchone()
            if dup:
                errors["code"] = "Ya existe un producto con este código."
        if errors:
            raise ValidationError("Revise los datos del producto.", errors)
        return {"code": code, "description": description, "category_id": category_id, "unit": unit,
                "price_cents": price, "discount_price_cents": discount_price, "cost_cents": cost or 0,
                "tax_rate_bp": tax_bp, "min_stock_milli": min_milli,
                "active": int(bool(merged.get("active", True))),
                "external_ref": (str(merged.get("external_ref") or "").strip() or None)}

    def create(self, actor: Actor, data: dict) -> dict:
        actor.require("inventario.editar")
        initial = parse_decimal(data.get("initial_stock")) if data.get("initial_stock") not in (None, "") else None
        if initial is not None:
            actor.require("inventario.movimientos")
            if initial < 0:
                raise ValidationError("Revise los datos.", {"initial_stock": "La existencia inicial no puede ser negativa."})
        with self.db.tx() as conn:
            product_id = self.create_conn(conn, actor, data)
            if initial:
                product = conn.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
                self._entry(conn, actor, product, to_milli(initial), product["cost_cents"], "Existencia inicial",
                            "inventario_inicial", {}, None, kind="inicial")
        return self.get(actor, product_id)

    def create_conn(self, conn: sqlite3.Connection, actor: Actor, data: dict) -> int:
        values = self._validated(conn, data, None)
        now = iso()
        cols = ", ".join(values)
        product_id = conn.execute(
            f"INSERT INTO products({cols}, created_at, updated_at, created_by, updated_by)"
            f" VALUES ({', '.join('?' * len(values))}, ?, ?, ?, ?)", (*values.values(), now, now, actor.id, actor.id)).lastrowid
        record(conn, actor, "producto.crear", "product", product_id, {"codigo": values["code"],
                                                                      "descripcion": values["description"]})
        return product_id

    def update(self, actor: Actor, product_id: int, data: dict) -> dict:
        actor.require("inventario.editar")
        if "stock" in data or "stock_milli" in data:
            raise ValidationError("La existencia no se edita directamente: registre una entrada, salida o ajuste.")
        if "cost" in data and not actor.can("inventario.costos"):
            actor.require("inventario.costos")
        with self.db.tx() as conn:
            self.update_conn(conn, actor, product_id, data)
        return self.get(actor, product_id)

    def update_conn(self, conn: sqlite3.Connection, actor: Actor, product_id: int, data: dict) -> dict:
        current = conn.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
        if not current:
            raise NotFoundError("Producto no encontrado.")
        current = dict(current)
        if "version" in data and data["version"] is not None and int(data["version"]) != current["version"]:
            raise ConflictError("Otra persona modificó este producto. Recargue para ver los cambios.")
        values = self._validated(conn, data, current)
        changes = {k: {"antes": current[k], "despues": v} for k, v in values.items() if current[k] != v}
        if not changes:
            return {}
        sets = ", ".join(f"{k} = ?" for k in values)
        conn.execute(f"UPDATE products SET {sets}, updated_at = ?, updated_by = ?, version = version + 1 WHERE id = ?",
                     (*values.values(), iso(), actor.id, product_id))
        record(conn, actor, "producto.actualizar", "product", product_id, changes)
        return changes

    def set_active(self, actor: Actor, product_id: int, active: bool) -> dict:
        return self.update(actor, product_id, {"active": active})

    # ------------------------------------------------------------------ movimientos
    def apply_movement(self, conn: sqlite3.Connection, actor: Actor, product: sqlite3.Row | dict, kind: str, qty_milli: int,
                       reason: str, unit_cost_cents: int, reference: tuple[str, int] | None = None,
                       idempotency_key: str | None = None, allow_negative: bool | None = None) -> int:
        """Inserta un movimiento (la existencia cambia por trigger). No contabiliza: lo hace quien llama."""
        if qty_milli == 0:
            raise ValidationError("La cantidad no puede ser cero.")
        if not (reason or "").strip():
            raise ValidationError("Indique el motivo del movimiento.", {"reason": "El motivo es obligatorio."})
        stock = conn.execute("SELECT stock_milli FROM products WHERE id = ?", (product["id"],)).fetchone()[0]
        after = stock + qty_milli
        if allow_negative is None:
            allow_negative = bool(self.settings.get("inventario", conn)["allow_negative_stock"])
        if after < 0 and not allow_negative:
            raise ValidationError(
                f"Existencias insuficientes de {product['code']}: hay {from_milli(stock)} y se requieren {from_milli(-qty_milli)}.",
                {"qty": f"Disponible: {from_milli(stock)}"})
        ref_type, ref_id = reference or (None, None)
        return conn.execute(
            "INSERT INTO inventory_movements(product_id, kind, qty_milli, stock_after_milli, unit_cost_cents, reason,"
            " reference_type, reference_id, user_id, created_at, idempotency_key) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (product["id"], kind, qty_milli, after, unit_cost_cents, reason.strip()[:300], ref_type, ref_id, actor.id,
             iso(), idempotency_key)).lastrowid

    def register_movement(self, actor: Actor, product_id: int, data: dict, idempotency_key: str) -> dict:
        actor.require("inventario.movimientos")
        kind = data.get("kind")
        if kind not in ("entrada", "salida", "ajuste"):
            raise ValidationError("Tipo de movimiento no válido.", {"kind": "Elija entrada, salida o ajuste."})
        if not idempotency_key or len(idempotency_key) > 80:
            raise ValidationError("Falta la clave de idempotencia de la operación.")
        reason = (data.get("reason") or "").strip()
        if not reason:
            raise ValidationError("Indique el motivo.", {"reason": "El motivo es obligatorio."})
        with self.db.tx() as conn:
            done = conn.execute("SELECT id FROM inventory_movements WHERE idempotency_key = ?", (idempotency_key,)).fetchone()
            if done:
                return {"duplicate": True, "movement_id": done[0], "product": self._product_conn(conn, product_id)}
            product = conn.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
            if not product:
                raise NotFoundError("Producto no encontrado.")
            if not product["active"]:
                raise ValidationError("El producto está desactivado. Actívelo antes de mover existencias.")
            if kind == "ajuste":
                target = parse_decimal(data.get("counted_qty"))
                if target is None or target < 0:
                    raise ValidationError("Revise la cantidad.", {"counted_qty": "Escriba la cantidad contada (cero o más)."})
                delta = to_milli(target) - product["stock_milli"]
                if delta == 0:
                    raise ValidationError("La cantidad contada coincide con la existencia: no hay nada que ajustar.",
                                          {"counted_qty": "Sin diferencia."})
                move_id = self._adjust(conn, actor, product, delta, reason, idempotency_key)
            else:
                qty = parse_decimal(data.get("qty"))
                if qty is None or qty <= 0:
                    raise ValidationError("Revise la cantidad.", {"qty": "La cantidad debe ser mayor que cero."})
                qty_milli = to_milli(qty)
                if product["unit"] == "und" and qty_milli % 1000:
                    raise ValidationError("Revise la cantidad.", {"qty": "Este producto se maneja en unidades enteras."})
                counterpart = data.get("counterpart") or ("compra_contado" if kind == "entrada" else "merma")
                if kind == "entrada":
                    if counterpart not in ENTRY_COUNTERPARTS:
                        raise ValidationError("Origen no válido.", {"counterpart": "Elija el origen de la entrada."})
                    raw_cost = data.get("unit_cost")
                    cost = to_cents(raw_cost) if raw_cost not in (None, "") else product["cost_cents"]
                    if cost < 0:
                        raise ValidationError("Revise el costo.", {"unit_cost": "El costo no puede ser negativo."})
                    move_id = self._entry(conn, actor, product, qty_milli, cost, reason, counterpart, data, idempotency_key)
                else:
                    if counterpart not in EXIT_COUNTERPARTS:
                        raise ValidationError("Destino no válido.", {"counterpart": "Elija el destino de la salida."})
                    move_id = self._exit(conn, actor, product, qty_milli, reason, counterpart, idempotency_key)
            return {"duplicate": False, "movement_id": move_id, "product": self._product_conn(conn, product_id)}

    def _entry(self, conn: sqlite3.Connection, actor: Actor, product, qty_milli: int, unit_cost: int, reason: str,
               counterpart: str, data: dict, idempotency_key: str | None, kind: str = "entrada") -> int:
        stock = product["stock_milli"]
        # Costo promedio ponderado; con existencia negativa o cero, el costo de la entrada.
        if stock > 0 and product["cost_cents"]:
            new_cost = round_cents((Decimal(stock) * product["cost_cents"] + Decimal(qty_milli) * unit_cost) / (stock + qty_milli))
        else:
            new_cost = unit_cost
        move_id = self.apply_movement(conn, actor, product, kind, qty_milli, reason, unit_cost,
                                      idempotency_key=idempotency_key)
        if new_cost != product["cost_cents"]:
            conn.execute("UPDATE products SET cost_cents = ?, updated_at = ?, updated_by = ?, version = version + 1 WHERE id = ?",
                         (new_cost, iso(), actor.id, product["id"]))
        value = round_cents(Decimal(qty_milli) * unit_cost / 1000)
        today = local_today(self._tz())
        desc = f"{ENTRY_COUNTERPARTS[counterpart]}: {product['code']} × {from_milli(qty_milli)}"
        if counterpart == "compra_credito":
            try:
                supplier_id = int(data.get("supplier_id"))
            except (TypeError, ValueError):
                raise ValidationError("Faltan datos.", {"supplier_id": "Seleccione el proveedor."}) from None
            if value <= 0:
                raise ValidationError("Revise el costo.", {"unit_cost": "Una compra a crédito necesita costo unitario."})
            due = parse_date(data.get("due_date"), "Vencimiento") if data.get("due_date") else today + timedelta(days=30)
            self.payables.create_conn(conn, actor, supplier_id, value, today, due, desc,
                                      number=(data.get("invoice_number") or "").strip() or None,
                                      source_type="entrada_inventario", source_id=move_id, account=ledger.INVENTORY)
        elif value > 0:
            if counterpart == "compra_contado":
                method, account = ledger.method_account(conn, int(data.get("method_id") or 0))
                if account is None:
                    raise ValidationError("Método no válido.", {"method_id": "Para compras a crédito elija «Compra a crédito»."})
                credit = Line(account, credit=value, memo=method["name"])
            elif counterpart == "inventario_inicial":
                credit = Line(ledger.CAPITAL, credit=value)
            else:
                credit = Line(ledger.OTHER_INCOME, credit=value)
            ledger.post(conn, actor, today, desc, "movimiento", move_id, [Line(ledger.INVENTORY, debit=value), credit])
        record(conn, actor, f"inventario.{kind}", "product", product["id"],
               {"cantidad": str(from_milli(qty_milli)), "costo_unitario": str(from_cents(unit_cost)), "origen": counterpart,
                "motivo": reason, "movimiento": move_id})
        return move_id

    def _exit(self, conn: sqlite3.Connection, actor: Actor, product, qty_milli: int, reason: str, counterpart: str,
              idempotency_key: str | None) -> int:
        move_id = self.apply_movement(conn, actor, product, "salida", -qty_milli, reason, product["cost_cents"],
                                      idempotency_key=idempotency_key)
        value = round_cents(Decimal(qty_milli) * product["cost_cents"] / 1000)
        expense = ledger.INVENTORY_LOSS if counterpart == "merma" else "5195"
        ledger.post(conn, actor, local_today(self._tz()), f"{EXIT_COUNTERPARTS[counterpart]}: {product['code']} ×"
                    f" {from_milli(qty_milli)}", "movimiento", move_id,
                    [Line(expense, debit=value, memo=reason), Line(ledger.INVENTORY, credit=value)])
        record(conn, actor, "inventario.salida", "product", product["id"],
               {"cantidad": str(from_milli(qty_milli)), "destino": counterpart, "motivo": reason, "movimiento": move_id})
        return move_id

    def _adjust(self, conn: sqlite3.Connection, actor: Actor, product, delta: int, reason: str,
                idempotency_key: str | None, kind: str = "ajuste") -> int:
        move_id = self.apply_movement(conn, actor, product, kind, delta, reason, product["cost_cents"],
                                      idempotency_key=idempotency_key, allow_negative=True)
        value = round_cents(Decimal(abs(delta)) * product["cost_cents"] / 1000)
        lines = ([Line(ledger.INVENTORY, debit=value), Line(ledger.OTHER_INCOME, credit=value)] if delta > 0 else
                 [Line(ledger.INVENTORY_LOSS, debit=value), Line(ledger.INVENTORY, credit=value)])
        ledger.post(conn, actor, local_today(self._tz()), f"{'Sobrante' if delta > 0 else 'Faltante'} por {kind}:"
                    f" {product['code']} {from_milli(delta)}", "movimiento", move_id, lines)
        record(conn, actor, f"inventario.{kind}", "product", product["id"],
               {"diferencia": str(from_milli(delta)), "motivo": reason, "movimiento": move_id})
        return move_id

    def _product_conn(self, conn: sqlite3.Connection, product_id: int) -> dict:
        return dict(conn.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone())

    # ------------------------------------------------------------------ exportación
    def export(self, actor: Actor, fmt: str = "xlsx") -> tuple[bytes, str, str]:
        actor.require("inventario.importar")
        include_cost = actor.can("inventario.costos")
        with self.db.read() as conn:
            rows = conn.execute("SELECT p.*, c.name AS category FROM products p LEFT JOIN categories c ON c.id = p.category_id"
                                " ORDER BY p.code COLLATE NOCASE").fetchall()
        header = ["Código", "Descripción", "Categoría", "Unidad", "Precio sin descuento", "Precio con descuento"]
        header += ["Costo"] if include_cost else []
        header += ["IVA %", "Cantidad disponible", "Existencia mínima", "Activo"]
        data = []
        for r in rows:
            line = [r["code"], r["description"], r["category"] or "", r["unit"], from_cents(r["price_cents"]),
                    from_cents(r["discount_price_cents"]) if r["discount_price_cents"] is not None else None]
            line += [from_cents(r["cost_cents"])] if include_cost else []
            line += [Decimal(r["tax_rate_bp"]) / 100, from_milli(r["stock_milli"]), from_milli(r["min_stock_milli"]),
                     "Sí" if r["active"] else "No"]
            data.append(line)
        stamp = local_today(self._tz()).isoformat()
        with self.db.tx() as conn:
            record(conn, actor, "inventario.exportar", "product", None, {"formato": fmt, "filas": len(data)})
        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=";")
            writer.writerow(header)
            for line in data:
                writer.writerow(["" if v is None else str(v).replace(".", ",") if isinstance(v, Decimal) else v for v in line])
            return ("﻿" + buf.getvalue()).encode("utf-8"), f"inventario-{stamp}.csv", "text/csv; charset=utf-8"
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventario"
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2937")
        for line in data:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in line])
        for col, width in zip("ABCDEFGHIJK", (16, 44, 18, 9, 18, 18, 14, 9, 18, 16, 8), strict=False):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        return (out.getvalue(), f"inventario-{stamp}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
