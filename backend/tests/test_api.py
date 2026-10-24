"""API de extremo a extremo (mismo contrato que usa la interfaz React)."""
from __future__ import annotations

import io
import json

import openpyxl

from .conftest import key, login


def test_health_headers_and_spa_without_build(client):
    r = client.get("/api/health")
    assert r.json()["status"] == "ok"
    assert r.headers["x-content-type-options"] == "nosniff" and "frame-ancestors 'none'" in r.headers["content-security-policy"]
    assert client.get("/inventario").status_code == 503
    assert client.get("/api/no-existe").status_code == 404


def test_full_sale_flow_over_http(client, ctx):
    login(client)
    r = client.post("/api/products", json={"code": "TAL-01", "description": "Taladro", "price": "350000.50",
                                           "discount_price": "320000", "cost": "200000", "tax_rate": "19",
                                           "initial_stock": "3"})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["price"] == "350000.50" and p["stock"] == "3" and p["tax_rate"] == "19"   # texto decimal exacto
    bad = client.post("/api/products", json={"code": "TAL-01", "description": "", "price": "x"})
    assert bad.status_code == 422 and {"code", "description", "price"} <= set(bad.json()["fields"])
    methods = {m["code"]: m["id"] for m in client.get("/api/payment-methods").json()}
    quote = client.post("/api/sales/quote", json={"lines": [{"product_id": p["id"], "qty": "2"}]}).json()
    assert quote["totals"]["total"] == "640000.00" and quote["can_sell"]
    body = {"lines": [{"product_id": p["id"], "qty": "2"}], "payments": [{"method_id": methods["efectivo"], "amount": "640000"}]}
    assert client.post("/api/sales", json=body).status_code == 422          # sin Idempotency-Key
    k = key()
    first = client.post("/api/sales", json=body, headers={"Idempotency-Key": k}).json()
    again = client.post("/api/sales", json=body, headers={"Idempotency-Key": k}).json()
    assert first["number"] == again["number"] and again["duplicate"]
    assert client.get(f"/api/products/{p['id']}").json()["stock"] == "1"
    over = client.post("/api/sales", json=body, headers={"Idempotency-Key": key()})
    assert over.status_code == 422 and "insuficientes" in over.json()["detail"]
    receipt = client.get(f"/api/sales/{first['id']}/receipt").json()
    assert first["number"] in receipt["text"]
    history = client.get("/api/sales").json()
    assert history["total"] == 1 and history["items"][0]["total"] == "640000.00"


def test_documents_over_http(client, ctx):
    login(client)
    p = client.post("/api/products", json={"code": "C1", "description": "Cable", "price": "1000", "initial_stock": "5"}).json()
    types = {t["code"]: t for t in client.get("/api/document-types").json()}
    t = types["cotizacion"]
    sample = client.get(f"/api/document-types/{t['id']}/sample-template")
    assert sample.status_code == 200
    up = client.post(f"/api/document-types/{t['id']}/template",
                     files={"file": ("cot.xlsx", sample.content, "application/octet-stream")})
    assert up.status_code == 200 and up.json()["template_problems"] == []
    doc = client.post("/api/documents", json={"type_id": t["id"], "customer_name": "Cliente",
                                              "lines": [{"product_id": p["id"], "qty": "3"}]}).json()
    assert doc["total"] == "3000.00"
    preview = client.get(f"/api/documents/{doc['id']}/preview").json()
    assert "Cable" in preview["html"]
    f = client.get(f"/api/documents/{doc['id']}/file")
    ws = openpyxl.load_workbook(io.BytesIO(f.content)).active
    assert any(c.value == "Cable" for row in ws.iter_rows() for c in row)


def test_import_over_http(client):
    login(client)
    wb = openpyxl.Workbook()
    wb.active.append(["Código", "Descripción", "Precio", "Cantidad"])
    wb.active.append(["IMP1", "Importado", "2500", "4"])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("inv.xlsx", buf.getvalue(), "application/octet-stream")}
    info = client.post("/api/imports/inspect", files=files).json()
    run = client.post("/api/imports/preview", files=files, data={"options": json.dumps({"columns": info["suggested_columns"]})}).json()
    assert run["summary"]["counts"]["new"] == 1
    applied = client.post(f"/api/imports/{run['id']}/apply").json()
    assert applied["status"] == "aplicada"
    assert client.get("/api/products/lookup?q=IMP1").json()[0]["stock"] == "4"


def test_finance_and_management_over_http(client):
    login(client)
    r = client.post("/api/finance/entries", json={"description": "Capital", "lines": [
        {"account_code": "1105", "debit": "500000"}, {"account_code": "3105", "credit": "500000"}]},
        headers={"Idempotency-Key": key()})
    assert r.status_code == 200
    bs = client.get("/api/finance/balance-sheet").json()
    assert bs["current"]["balanced"] and bs["current"]["total_assets"] == "500000.00"
    assert client.get("/api/finance/ar").status_code == 200
    assert client.get("/api/finance/ap/aging").status_code == 200
    assert client.post("/api/finance/break-even", json={"fixed_costs": "100", "price": "10", "variable_cost": "5"}).json()[
        "break_even_units"] == "20.00"
    users = client.post("/api/users", json={"username": "cajero", "full_name": "Cajero", "password": "Cajero clave 1",
                                             "permissions": ["ventas.vender"]})
    assert users.status_code == 422          # la contraseña contiene el usuario
    created = client.post("/api/users", json={"username": "cajero", "full_name": "Cajero", "password": "Mostrador 2026",
                                               "permissions": ["ventas.vender"]}).json()
    assert created["permissions"] == ["ventas.vender"]
    assert client.put("/api/settings/empresa", json={"name": "Mi Ferretería"}).json()["name"] == "Mi Ferretería"
    assert client.get("/api/company/public").json()["name"] == "Mi Ferretería"
    backup = client.post("/api/backups").json()
    assert client.get(f"/api/backups/{backup['file']}/verify").json()["ok"]
    restore = client.post("/api/backups/restore", json={"file": backup["file"], "confirm": "no"})
    assert restore.status_code == 422          # exige confirmación explícita


def test_setup_flow_over_http(tmp_path):
    from fastapi.testclient import TestClient

    from pospy.api.app import create_app
    from pospy.config import Config
    from pospy.context import AppContext
    app = AppContext(Config.from_env(data_dir=tmp_path, master_key="x"))
    c = TestClient(create_app(app))
    c.headers.update({"X-Requested-With": "pos-py"})
    status = c.get("/api/setup/status").json()
    assert status["needs_setup"] and status["code_file"].endswith("setup-code.txt")
    code = app.auth.setup_file.read_text().strip()
    assert c.post("/api/setup", json={"code": "X", "username": "jefa", "full_name": "Jefa", "password": "Clave buena 2026"}
                  ).status_code == 422
    assert c.post("/api/setup", json={"code": code, "username": "jefa", "full_name": "Jefa", "password": "Clave buena 2026"}
                  ).status_code == 200
    assert not c.get("/api/setup/status").json()["needs_setup"]
    assert login(c, "jefa", "Clave buena 2026")["user"]["role"] == "admin"
