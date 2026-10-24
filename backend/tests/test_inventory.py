"""3. Creación y edición de productos · 4. Entradas y salidas · 12. Validación · integridad."""
from __future__ import annotations

import io
import sqlite3

import openpyxl
import pytest

from pospy.core.errors import ConflictError, PermissionDenied, ValidationError

from .conftest import balance, key


def test_create_and_edit_product(ctx, admin, product):
    assert product["stock_milli"] == 10000 and product["cost_cents"] == 1000000
    updated = ctx.inventory.update(admin, product["id"], {"description": "Martillo de uña 16 oz", "price": "21000",
                                                          "version": product["version"]})
    assert updated["price_cents"] == 2100000 and updated["stock_milli"] == 10000
    with pytest.raises(ConflictError):
        ctx.inventory.update(admin, product["id"], {"price": "1", "version": product["version"]})
    with ctx.db.read() as conn:
        actions = [r[0] for r in conn.execute("SELECT action FROM audit_log WHERE entity = 'product'")]
    assert "producto.crear" in actions and "producto.actualizar" in actions


def test_product_validation_messages_by_field(ctx, admin, product):
    with pytest.raises(ValidationError) as exc:
        ctx.inventory.create(admin, {"code": "MART-01", "description": "", "price": "-5", "discount_price": "abc"})
    fields = exc.value.fields
    assert {"code", "description", "price", "discount_price"} <= set(fields)
    with pytest.raises(ValidationError) as exc:
        ctx.inventory.create(admin, {"code": "X1", "description": "X", "price": "100", "discount_price": "150"})
    assert "discount_price" in exc.value.fields


def test_stock_cannot_be_edited_directly(ctx, admin, product):
    with pytest.raises(ValidationError):
        ctx.inventory.update(admin, product["id"], {"stock": "99"})
    with ctx.db.tx() as conn, pytest.raises(sqlite3.IntegrityError):
        conn.execute("UPDATE products SET stock_milli = 1 WHERE id = ?", (product["id"],))
    with ctx.db.tx() as conn, pytest.raises(sqlite3.IntegrityError):
        conn.execute("DELETE FROM inventory_movements")


def test_entry_exit_adjust_history_and_average_cost(ctx, admin, product, cash_id):
    ctx.inventory.register_movement(admin, product["id"], {"kind": "entrada", "qty": "10", "unit_cost": "12000",
                                                           "reason": "Compra", "counterpart": "compra_contado",
                                                           "method_id": cash_id}, key())
    p = ctx.inventory.get(admin, product["id"])
    assert p["stock_milli"] == 20000 and p["cost_cents"] == 1100000   # (10×10.000 + 10×12.000) / 20
    ctx.inventory.register_movement(admin, product["id"], {"kind": "salida", "qty": "2", "reason": "Dañados",
                                                           "counterpart": "merma"}, key())
    ctx.inventory.register_movement(admin, product["id"], {"kind": "ajuste", "counted_qty": "17", "reason": "Conteo"}, key())
    p = ctx.inventory.get(admin, product["id"])
    assert p["stock_milli"] == 17000
    history = ctx.inventory.history(admin, product["id"])["items"]
    assert [m["kind"] for m in history] == ["ajuste", "salida", "entrada", "inicial"]
    assert all(m["username"] == "admin" and m["reason"] for m in history)
    assert history[0]["qty_milli"] == -1000 and history[0]["stock_after_milli"] == 17000
    # Inventario contable = valorizado: 17 × 11.000
    assert balance(ctx, "1435") == 17 * 1100000
    assert balance(ctx, "1105") == -10 * 1200000


def test_exit_cannot_leave_negative_stock(ctx, admin, product):
    with pytest.raises(ValidationError, match="insuficientes"):
        ctx.inventory.register_movement(admin, product["id"], {"kind": "salida", "qty": "11", "reason": "x",
                                                               "counterpart": "merma"}, key())
    assert ctx.inventory.get(admin, product["id"])["stock_milli"] == 10000


def test_movement_idempotency(ctx, admin, product):
    k = key()
    data = {"kind": "salida", "qty": "1", "reason": "Muestra", "counterpart": "consumo_interno"}
    first = ctx.inventory.register_movement(admin, product["id"], data, k)
    second = ctx.inventory.register_movement(admin, product["id"], data, k)
    assert second["duplicate"] and second["movement_id"] == first["movement_id"]
    assert ctx.inventory.get(admin, product["id"])["stock_milli"] == 9000


def test_credit_purchase_creates_payable(ctx, admin, product):
    supplier = ctx.catalog.save_party(admin, "proveedores", {"name": "Ferretería Mayorista", "identification": "900"})
    ctx.inventory.register_movement(admin, product["id"], {
        "kind": "entrada", "qty": "5", "unit_cost": "9000", "reason": "Factura FM-1", "counterpart": "compra_credito",
        "supplier_id": supplier["id"], "invoice_number": "FM-1", "due_date": "2099-01-01"}, key())
    items = ctx.payables.list(admin)
    assert len(items) == 1 and items[0]["balance_cents"] == 4500000
    assert balance(ctx, "2205") == -4500000


def test_deactivate_product_and_employee_restrictions(ctx, admin, employee, product):
    ctx.inventory.set_active(admin, product["id"], False)
    assert ctx.inventory.list(admin, status="activos")["total"] == 0
    assert ctx.inventory.list(admin, status="inactivos")["total"] == 1
    with pytest.raises(PermissionDenied):
        ctx.inventory.create(employee, {"code": "B", "description": "B", "price": "1"})
    listed = ctx.inventory.list(employee, status="todos")["items"][0]
    assert "cost_cents" not in listed   # sin permiso de costos


def test_search_sort_paginate_and_low_stock(ctx, admin):
    for i in range(30):
        ctx.inventory.create(admin, {"code": f"P{i:02d}", "description": f"Tornillo {i}", "price": str(100 + i),
                                     "min_stock": "5", "initial_stock": str(i)})
    page = ctx.inventory.list(admin, q="tornillo", sort="price", direction="desc", page=2, page_size=10)
    assert page["total"] == 30 and page["pages"] == 3 and page["items"][0]["code"] == "P19"
    low = ctx.inventory.list(admin, low_stock=True, page_size=100)
    assert low["total"] == 6   # existencias 0..5 con mínimo 5
    assert ctx.inventory.lookup(admin, "P07")[0]["code"] == "P07"


def test_export_xlsx_and_csv(ctx, admin, product):
    data, name, _ = ctx.inventory.export(admin, "xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value == "Código" and ws["A2"].value == "MART-01" and name.endswith(".xlsx")
    data, name, _ = ctx.inventory.export(admin, "csv")
    text = data.decode("utf-8-sig")
    assert "MART-01;Martillo de uña;" in text and "20000,00" in text


def test_ledger_invariant_stock_equals_movements(ctx, admin, product):
    result = ctx.finance.integrity(admin)
    assert result["ok"], result
