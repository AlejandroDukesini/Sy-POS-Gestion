"""7. Cotizaciones y órdenes · 8. Conservación del formato de plantillas · 9. Configuración e impresión."""
from __future__ import annotations

import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Border, Font, PatternFill, Side

from pospy.core.errors import DomainError, PermissionDenied, ValidationError
from pospy.documents import com_excel
from pospy.documents.templates import TemplateMappingError
from pospy.printing import escpos, printers

from .conftest import key


def make_template(path, rows=5, chart=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotizacion"
    ws["A1"] = "{{empresa_nombre}}"
    ws["A1"].font = Font(bold=True, size=18, color="1D4ED8")
    ws.merge_cells("A1:D1")
    ws["E1"] = "Cotización N° {{numero}}"
    ws["E2"] = "{{fecha}}"
    ws["A3"] = "Cliente: {{cliente_nombre}}"
    ws["A4"] = "Placa: {{extra.placa}}"
    thin = Side(style="thin", color="000000")
    for col, fld in zip("ABCDEF", ["n", "codigo", "descripcion", "cantidad", "precio", "total"], strict=True):
        ws[f"{col}6"] = col
        ws[f"{col}6"].fill = PatternFill("solid", fgColor="FFCC00")
        ws[f"{col}7"] = "{{linea.%s}}" % fld
        for r in range(7, 7 + rows):
            ws[f"{col}{r}"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
            ws[f"F{r}"].number_format = '"$"#,##0'
    ws[f"E{7 + rows}"] = "Suma:"
    ws[f"F{7 + rows}"] = f"=SUM(F7:F{6 + rows})"
    ws[f"F{8 + rows}"] = "{{total}}"
    ws.column_dimensions["C"].width = 42
    ws.oddFooter.center.text = "Pie de página de la empresa"
    if chart:
        c = BarChart()
        c.add_data(Reference(ws, min_col=6, min_row=7, max_row=6 + rows))
        ws.add_chart(c, "H2")
    wb.save(path)
    return path


@pytest.fixture()
def quote_type(ctx, admin, tmp_path):
    types = {t["code"]: t for t in ctx.documents.types(admin)}
    assert {"cotizacion", "orden"} <= set(types)
    t = types["cotizacion"]
    ctx.documents.save_type(admin, {"name": "Cotización", "prefix": "COT-", "validity_days": 15, "printer": "DIGITALPOS",
                                    "fields": [{"label": "Placa", "key": "placa", "required": True}],
                                    "page_setup": {"paper": "letter", "orientation": "landscape", "margins_cm": 1,
                                                   "fit_width": True, "scale": 100}}, t["id"])
    path = make_template(tmp_path / "plantilla.xlsx")
    result = ctx.documents.upload_template(admin, t["id"], "plantilla.xlsx", path.read_bytes())
    assert result["template_problems"] == [], result["template_problems"]
    return result


def new_doc(ctx, actor, quote_type, product, qty="2", **extra):
    return ctx.documents.create(actor, {"type_id": quote_type["id"], "customer_name": "Taller Ruiz", "extra": {"placa": "ABC123"},
                                        "lines": [{"product_id": product["id"], "qty": qty}], **extra}, key())


def test_unlimited_document_types_and_stats(ctx, admin):
    created = ctx.documents.save_type(admin, {"name": "Remisión", "prefix": "REM-", "page_setup": {}})
    assert created["code"] == "remision" and created["documents"] == 0
    assert len(ctx.documents.types(admin)) == 3


def test_create_quote_with_numbering_extra_fields_and_totals(ctx, admin, employee, quote_type, product):
    with pytest.raises(ValidationError, match="Faltan datos"):
        ctx.documents.create(admin, {"type_id": quote_type["id"], "customer_name": "X",
                                     "lines": [{"product_id": product["id"], "qty": "1"}]})
    doc = new_doc(ctx, employee, quote_type, product)
    assert doc["number"] == "COT-000001" and doc["total_cents"] == 3600000 and doc["valid_until"]
    assert new_doc(ctx, admin, quote_type, product)["number"] == "COT-000002"
    assert ctx.inventory.get(admin, product["id"])["stock_milli"] == 10000   # un documento no mueve inventario
    stats = ctx.documents.get_type(admin, quote_type["id"])
    assert stats["documents"] == 2 and stats["last_30_days"] == 2


def test_generation_preserves_template_format(ctx, admin, quote_type, product, tmp_path):
    doc = new_doc(ctx, admin, quote_type, product, qty="3")
    out = ctx.documents.generate(admin, doc["id"])
    ws = openpyxl.load_workbook(out)["Cotizacion"]
    assert ws["A1"].value == "Mi empresa" and ws["A1"].font.b and ws["A1"].font.sz == 18
    assert "A1:D1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["E1"].value == f"Cotización N° {doc['number']}"
    assert ws["A4"].value == "Placa: ABC123"
    assert ws["C7"].value == "Martillo de uña" and ws["D7"].value == 3 and ws["F7"].value == 54000
    assert ws["F7"].number_format == '"$"#,##0' and ws["C8"].border.left.style == "thin"
    assert ws["A6"].fill.fgColor.rgb.endswith("FFCC00")
    assert ws["C8"].value is None                          # filas libres vaciadas
    assert ws["F12"].value == "=SUM(F7:F11)"               # la fórmula de la plantilla se conserva
    assert ws["F13"].value == 54000
    assert ws.column_dimensions["C"].width == 42
    assert ws.oddFooter.center.text == "Pie de página de la empresa"
    assert ws.page_setup.orientation == "landscape"
    template = ctx.documents.templates_dir / quote_type["template_file"]
    assert openpyxl.load_workbook(template)["Cotizacion"]["C7"].value == "{{linea.descripcion}}"   # original intacta
    html = ctx.documents.preview(admin, doc["id"])["html"]
    assert "Martillo de uña" in html and "colspan" in html


def test_capacity_and_template_drift_are_detected(ctx, admin, quote_type, product, tmp_path):
    many = [ctx.inventory.create(admin, {"code": f"X{i}", "description": f"Item {i}", "price": "10", "initial_stock": "1"})
            for i in range(6)]
    doc = ctx.documents.create(admin, {"type_id": quote_type["id"], "customer_name": "C", "extra": {"placa": "Z"},
                                       "lines": [{"product_id": p["id"], "qty": "1"} for p in many]})
    with pytest.raises(TemplateMappingError, match="admite 5"):
        ctx.documents.generate(admin, doc["id"])
    template = ctx.documents.templates_dir / quote_type["template_file"]
    wb = openpyxl.load_workbook(template)
    wb["Cotizacion"]["E1"] = "Sin marcador"
    wb.save(template)
    with pytest.raises(TemplateMappingError, match="numero"):
        ctx.documents.generate(admin, new_doc(ctx, admin, quote_type, product)["id"])


def test_manual_field_assignment(ctx, admin, quote_type, product):
    t = ctx.documents.assign_field(admin, quote_type["id"], "cliente_telefono", "E3")
    assert t["mapping"]["fields"]["cliente_telefono"] == "E3" and not t["template_problems"]
    doc = new_doc(ctx, admin, quote_type, product, customer_phone="300 123")
    ws = openpyxl.load_workbook(ctx.documents.generate(admin, doc["id"]))["Cotizacion"]
    assert ws["E3"].value == "300 123"


def test_sample_template_is_valid(ctx, admin, tmp_path):
    t = next(x for x in ctx.documents.types(admin) if x["code"] == "orden")
    data, name = ctx.documents.sample_template(admin, t["id"])
    result = ctx.documents.upload_template(admin, t["id"], name, data)
    assert result["template_problems"] == [] and result["mapping"]["capacity"] == 20


def test_print_without_printer_fails_clearly_and_is_logged(ctx, admin, quote_type, product, monkeypatch):
    doc = new_doc(ctx, admin, quote_type, product)
    monkeypatch.setattr(com_excel, "com_available", lambda: False)
    with pytest.raises(DomainError, match="Microsoft Excel"):
        ctx.documents.print(admin, doc["id"])
    jobs = ctx.printing.jobs(admin, "documento", doc["id"])
    assert jobs[0]["status"] == "error" and jobs[0]["printer"] == "DIGITALPOS"


def test_print_document_uses_configured_printer(ctx, admin, quote_type, product, monkeypatch):
    calls = []
    monkeypatch.setattr(com_excel, "print_workbook", lambda path, printer, copies, setup: calls.append((printer, copies, setup)))
    doc = new_doc(ctx, admin, quote_type, product)
    assert ctx.documents.print(admin, doc["id"])["printer"] == "DIGITALPOS"
    ctx.documents.print(admin, doc["id"], copies=2)   # reimpresión
    assert calls[0][0] == "DIGITALPOS" and calls[0][2]["orientation"] == "landscape" and calls[1][1] == 2
    detail = ctx.documents.get(admin, doc["id"])
    assert len(detail["prints"]) == 2 and all(p["status"] == "ok" for p in detail["prints"])


def test_employee_needs_print_permission(ctx, admin, employee, quote_type, product):
    doc = new_doc(ctx, admin, quote_type, product)
    ctx.users.update(admin, employee.id, {"permissions": ["documentos.ver"]})
    from .conftest import EMPLOYEE_PASSWORD
    limited = ctx.auth.login("empleado", EMPLOYEE_PASSWORD).actor
    with pytest.raises(PermissionDenied):
        ctx.documents.print(limited, doc["id"])
    with pytest.raises(PermissionDenied):
        ctx.documents.save_type(limited, {"name": "X"})


@pytest.mark.skipif(not com_excel.com_available(), reason="requiere Microsoft Excel")
def test_template_with_chart_is_filled_with_excel_and_keeps_chart(ctx, admin, product, tmp_path):
    t = next(x for x in ctx.documents.types(admin) if x["code"] == "orden")
    ctx.documents.save_type(admin, {"name": "Orden", "prefix": "ORD-", "fields": [{"label": "Placa", "key": "placa"}],
                                    "page_setup": {}}, t["id"])
    path = make_template(tmp_path / "con-grafico.xlsx", chart=True)
    t = ctx.documents.upload_template(admin, t["id"], "con-grafico.xlsx", path.read_bytes())
    assert "gráficos" in t["template_needs_excel"]
    doc = new_doc(ctx, admin, t, product)
    out = ctx.documents.generate(admin, doc["id"])
    with zipfile.ZipFile(out) as zf:
        assert any(n.startswith("xl/charts/") for n in zf.namelist())
    assert openpyxl.load_workbook(out)["Cotizacion"]["C7"].value == "Martillo de uña"


# ------------------------------------------------------------------ comprobante de venta (DIGITALPOS)
def test_escpos_receipt_content(ctx, admin, product, cash_id):
    ctx.settings.update("empresa", {"name": "Ferretería Ñandú", "tax_id": "900.123"}, admin)
    sale = ctx.sales.checkout(admin, {"lines": [{"product_id": product["id"], "qty": "2"}],
                                      "payments": [{"method_id": cash_id, "amount": "36000", "tendered": "40000"}]}, key())
    data = ctx.printing.receipt_bytes(ctx.sales.get(admin, sale["id"]))
    text = escpos.plain_text(data)
    assert "Ferretería Ñandú" in text and sale["number"] in text and "no es factura electrónica" in text
    assert "$ 36.000" in text and "Cambio" in text and "$ 4.000" in text
    assert data.endswith(escpos.CUT) and data.startswith(escpos.INIT)
    assert all(len(line) <= 48 for line in text.splitlines())


def test_receipt_printing_modes_and_errors(ctx, admin, product, cash_id, monkeypatch):
    sale = ctx.sales.checkout(admin, {"lines": [{"product_id": product["id"], "qty": "1"}],
                                      "payments": [{"method_id": cash_id, "amount": "18000"}]}, key())
    sent = []
    monkeypatch.setattr("pospy.services.printing.print_raw", lambda p, d, j: sent.append((p, d)))
    assert ctx.printing.print_receipt(admin, sale["id"])["printed"] and sent[0][0] == "DIGITALPOS"

    def fail(*_):
        raise printers.PrintError("Revise que esté encendida")
    monkeypatch.setattr("pospy.services.printing.print_raw", fail)
    with pytest.raises(printers.PrintError):
        ctx.printing.print_receipt(admin, sale["id"])
    ctx.settings.update("impresion", {"receipt_mode": "navegador"}, admin)
    assert ctx.printing.print_receipt(admin, sale["id"])["mode"] == "navegador"
    statuses = [j["status"] for j in ctx.printing.jobs(admin, "venta", sale["id"])]
    assert statuses == ["ok", "error", "ok"]


def test_find_printer_reports_available_printers(monkeypatch):
    monkeypatch.setattr(printers, "list_printers", lambda: {"available": True, "default": "",
                                                            "printers": [{"name": "DIGITALPOS 80"}, {"name": "HP"}], "reason": ""})
    assert printers.find_printer("digitalpos") == "DIGITALPOS 80"
    with pytest.raises(printers.PrintError, match="No se encontró"):
        printers.find_printer("EPSON")
